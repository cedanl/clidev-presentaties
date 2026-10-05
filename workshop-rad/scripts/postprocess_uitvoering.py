"""Nabewerking van het uitvoeringsdeck na powerclaude's build.mjs.

1. Titelslide: echte Npuls-achtergrond (Slide15.PNG: blauw met de bogen) en het volledige horizontale
   logo (stippenring + woordmerk) in plaats van een los beeldmerk.
2. "Zo loopt de middag": de tabel wordt een tijdlijn met blokken op schaal en per blok de bijbehorende skill.
3. Beeldmerk op de overige slides: powerclaude's indicatieve stippenster wordt de echte stippenring
   (zie fix_beeldmerk.py).

Gebruik (vanuit de root van clidev-presentaties):
    python workshop-rad/scripts/postprocess_uitvoering.py invoer.pptx uitvoer.pptx
"""
import copy
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

sys.path.insert(0, str(Path(__file__).parent))
import fix_beeldmerk  # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
ASSETS = ROOT / "workshop-rad/assets"
ACHTERGROND = ROOT / "public/npuls/powerpoint_slides/Slide15.PNG"

KLEUR = {"blauw": "3D68EC", "oranje": "DD784B", "geel": "F4D74B", "groen": "00AF81",
         "roze": "F4D9DC", "zwart": "000000", "wit": "FFFFFF"}
KOP, TEKST = "General Sans Semibold", "General Sans"


def rgb(naam):
    return RGBColor.from_string(KLEUR[naam])


def tekstvak(slide, x, y, w, h, tekst, grootte, kleur, font=TEKST, uitlijning=PP_ALIGN.LEFT,
             anker=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anker
    p = tf.paragraphs[0]
    p.alignment = uitlijning
    r = p.add_run()
    r.text = tekst
    r.font.size = Pt(grootte)
    r.font.name = font
    r.font.color.rgb = rgb(kleur)
    return tb


def vlak(slide, x, y, w, h, kleur, vorm=MSO_SHAPE.RECTANGLE, tekst=None, grootte=12, tekstkleur="zwart",
         font=KOP):
    sh = slide.shapes.add_shape(vorm, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(kleur)
    sh.line.fill.background()
    sh.shadow.inherit = False
    if tekst is not None:
        tf = sh.text_frame
        tf.margin_left = tf.margin_right = Inches(0.05)
        tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = tekst
        r.font.size = Pt(grootte)
        r.font.name = font
        r.font.color.rgb = rgb(tekstkleur)
    return sh


def naar_achter(slide, shape):
    tree = slide.shapes._spTree
    tree.remove(shape._element)
    tree.insert(2, shape._element)  # na nvGrpSpPr en grpSpPr


def is_stip(shape):
    geom = shape._element.spPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}prstGeom")
    return (geom is not None and geom.get("prst") == "ellipse"
            and shape.width <= int(0.12 * 914400) and shape.height <= int(0.12 * 914400))


def titelslide(prs):
    s = prs.slides[0]
    W, H = prs.slide_width, prs.slide_height
    for sh in [x for x in s.shapes if is_stip(x)]:  # indicatief beeldmerk
        sh._element.getparent().remove(sh._element)
    bg = s.shapes.add_picture(str(ACHTERGROND), 0, 0, W, H)
    bg.name = "Npuls achtergrond"
    naar_achter(s, bg)
    # Pill met "CEDA · DAIR-bijeenkomst" schuift rechts van het logo
    pill_vorm = [x for x in s.shapes if abs(x.top / 914400 - 0.68) < 0.02 and 3.0 < x.width / 914400 < 4.0]
    for sh in pill_vorm:
        sh.left = sh.left + Inches(2.55)
        sh.top = sh.top + Inches(0.02)
    logo = s.shapes.add_picture(str(ASSETS / "npuls-logo-horizontaal-wit.png"), Inches(0.6), Inches(0.5),
                                width=Inches(2.2))
    logo.name = "Npuls logo"
    # Pay-off linksboven onder het logo, zodat de bogen rechtsonder vrij blijven
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.startswith("Moving Education"):
            sh.left, sh.top = Inches(0.6), Inches(1.45)
            sh.width = Inches(4.0)
            sh.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT


BLOKKEN = [  # begin (min), duur, soort, naam, omschrijving, skill
    (0, 17, "geel", "Intro, data en casus", "CEDA, synthetische data en de opdracht", None),
    (17, 13, "geel", "Skills en devcontainer", "Wat is een skill, eerste prompt", None),
    (30, 22, "blauw", "Ronde 1: data verkennen", "In tweetallen met de agent", "/workshop-verkennen"),
    (52, 15, "oranje", "Spiegelmoment 1", "Kijk in de skill", "/workshop-verkennen"),
    (67, 22, "blauw", "Ronde 2: dashboard bouwen", "Wissel van rol", "/workshop-dashboard"),
    (89, 15, "oranje", "Spiegelmoment 2", "Kijk in de skill, onthulling", "/workshop-dashboard"),
    (104, 16, "groen", "Evaluatie en afsluiting", "Terugblik, peiling, rondje", "/workshop-reflectie"),
]
TEKST_OP = {"geel": "zwart", "oranje": "zwart", "blauw": "wit", "groen": "wit"}


def tijd(min_):
    return f"{min_ // 60}:{min_ % 60:02d}"


def tijdlijn(prs):
    s = next(sl for sl in prs.slides
             if any(sh.has_text_frame and sh.text_frame.text.strip() == "Zo loopt de middag" for sh in sl.shapes))
    for sh in [x for x in s.shapes if x.has_table]:
        sh._element.getparent().remove(sh._element)
    LINKS, BREEDTE, GAP = 0.6, 12.13, 0.05
    per_min = (BREEDTE - GAP * (len(BLOKKEN) - 1)) / 120
    BALK_Y, BALK_H = 2.3, 1.0
    x = LINKS
    for i, (begin, duur, soort, naam, omschr, skill) in enumerate(BLOKKEN):
        w = duur * per_min
        tekstvak(s, x, BALK_Y - 0.42, 1.0, 0.3, tijd(begin), 14, "blauw", KOP)
        vlak(s, x, BALK_Y, w, BALK_H, soort, tekst=f"{duur} min", grootte=18, tekstkleur=TEKST_OP[soort])
        # dunne lijn van balk naar beschrijving
        vlak(s, x, BALK_Y + BALK_H, w, 0.05, soort)
        tekstvak(s, x + 0.02, BALK_Y + BALK_H + 0.3, w - 0.1, 0.8, naam, 11, "zwart", KOP)
        tekstvak(s, x + 0.02, BALK_Y + BALK_H + 1.25, w - 0.1, 0.9, omschr, 12, "zwart")
        if skill:
            pil = vlak(s, x + 0.02, BALK_Y + BALK_H + 2.2, w - 0.04, 0.34, "wit", MSO_SHAPE.ROUNDED_RECTANGLE,
                       tekst=skill, grootte=9, tekstkleur="blauw")
            pil.text_frame.margin_left = pil.text_frame.margin_right = 0
            pil.text_frame.word_wrap = False
        x += w + GAP
    tekstvak(s, LINKS + BREEDTE - 1.0, BALK_Y - 0.42, 1.0, 0.3, "2:00", 14, "blauw", KOP, PP_ALIGN.RIGHT)
    # Legenda
    ly = 6.55
    lx = LINKS
    for soort, label in (("geel", "Uitleg"), ("blauw", "Werken in tweetallen"), ("oranje", "Spiegelmoment"),
                         ("groen", "Evaluatie en afsluiting")):
        vlak(s, lx, ly + 0.04, 0.22, 0.22, soort)
        t = tekstvak(s, lx + 0.32, ly, 2.4, 0.3, label, 12, "zwart", TEKST, anker=MSO_ANCHOR.MIDDLE)
        lx += 0.32 + 0.11 * len(label) + 0.5


def main(inp, uit):
    prs = Presentation(inp)
    titelslide(prs)
    tijdlijn(prs)
    prs.save(uit)
    fix_beeldmerk.main(uit)  # echte stippenring op de overige slides


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
