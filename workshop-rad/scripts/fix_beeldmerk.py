"""Vervangt het beeldmerk dat powerclaude tekent (een indicatieve, gevulde stippenster) door het
echte Npuls-beeldmerk (stippenring).

powerclaude tekent het merk rechtsboven als ~60 losse stippen uit een indicatief SVG. Dit script zoekt
die stippen per slide, verwijdert ze en plaatst op dezelfde plek het echte beeldmerk (PNG met
transparante achtergrond, wit of zwart naar gelang de stipkleur).

Gebruik (vanuit de root van clidev-presentaties), na het bouwen met build.mjs:
    python workshop-rad/scripts/fix_beeldmerk.py invoer.pptx [uitvoer.pptx]
Zonder uitvoer wordt het invoerbestand overschreven.
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

ROOT = Path(__file__).resolve().parents[2]
PNG = {"zwart": ROOT / "workshop-rad/assets/npuls-beeldmerk-zwart.png",
       "wit": ROOT / "workshop-rad/assets/npuls-beeldmerk-wit.png"}
MAX_STIP = int(0.12 * 914400)  # stippen zijn kleiner dan 0,12 inch
MIN_AANTAL = 20


def voorinstelling(shape):
    """prst-waarde uit de XML (python-pptx kent niet elke waarde, bijvoorbeeld 'line')."""
    geom = shape._element.spPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}prstGeom")
    return geom.get("prst") if geom is not None else None


def stipkleur(shape):
    try:
        rgb = str(shape.fill.fore_color.rgb)
    except Exception:
        return "zwart"
    r, g, b = (int(rgb[i:i + 2], 16) for i in (0, 2, 4))
    return "wit" if 0.299 * r + 0.587 * g + 0.114 * b > 140 else "zwart"


def main(inp, uit=None):
    prs = Presentation(inp)
    vervangen = 0
    for slide in prs.slides:
        stippen = [s for s in slide.shapes
                   if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and voorinstelling(s) == "ellipse"
                   and s.width <= MAX_STIP and s.height <= MAX_STIP]
        if len(stippen) < MIN_AANTAL:
            continue
        links = min(s.left for s in stippen); boven = min(s.top for s in stippen)
        rechts = max(s.left + s.width for s in stippen); onder = max(s.top + s.height for s in stippen)
        zijde = max(rechts - links, onder - boven)
        cx, cy = (links + rechts) // 2, (boven + onder) // 2
        kleur = stipkleur(stippen[0])
        for s in stippen:
            s._element.getparent().remove(s._element)
        slide.shapes.add_picture(str(PNG[kleur]), cx - zijde // 2, cy - zijde // 2, zijde, zijde)
        slide.shapes[-1].name = "Npuls beeldmerk"
        vervangen += 1
    prs.save(uit or inp)
    print(f"beeldmerk vervangen op {vervangen} slides -> {uit or inp}")


if __name__ == "__main__":
    main(*sys.argv[1:3])
